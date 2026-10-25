"""
Step 1 of the pipeline: turn a raw PDF/DOCX file into clean plain text.
This is the ONLY file that should know about file formats — everything
downstream just works with a string of text.
"""

import re
import pdfplumber
import docx


def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from a PDF, page by page, preserving reading order reasonably well."""
    text_chunks = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_chunks.append(page_text)
    return "\n".join(text_chunks)


def extract_text_from_docx(file_path: str) -> str:
    """
    Extract text from a DOCX file in document order. Includes table cells,
    because many resume templates lay out whole sections inside tables.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(file_path)
    lines = []

    def add_table(table):
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:      # merged cells repeat the same element
                    continue
                seen.add(id(cell._tc))
                for para in cell.paragraphs:
                    if para.text.strip():
                        lines.append(para.text)

    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            text = Paragraph(child, document).text
            if text.strip():
                lines.append(text)
        elif child.tag.endswith("}tbl"):
            add_table(Table(child, document))
    return "\n".join(lines)


def extract_text(file_path: str, file_extension: str) -> str:
    """Dispatch to the right extractor based on file type."""
    if file_extension == ".pdf":
        raw_text = extract_text_from_pdf(file_path)
    elif file_extension == ".docx":
        raw_text = extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_extension}")

    return clean_text(raw_text)


def clean_text(text: str) -> str:
    """Basic cleanup: normalize whitespace, strip odd bullet characters, drop empty lines."""
    text = re.sub(r"[•●▪·]", "-", text)          # normalize bullets
    text = text.replace("\u00a0", " ").replace("\u200b", "")   # non-breaking / zero-width spaces
    text = re.sub(r"[ \t]+", " ", text)            # collapse repeated spaces/tabs
    text = re.sub(r"\n{3,}", "\n\n", text)          # collapse excessive blank lines
    lines = [line.strip() for line in text.split("\n")]
    lines = [line for line in lines if line]        # drop empty lines
    return "\n".join(lines)


# --- Section segmentation ---

SECTION_HEADERS = {
    "summary": ["summary", "professional summary", "career summary", "objective", "career objective",
                "profile", "professional profile", "about me", "about"],
    "education": ["education", "academic background", "academic qualifications", "education and training",
                  "educational qualifications", "qualifications"],
    "experience": ["experience", "work experience", "professional experience", "employment history",
                   "work history", "employment", "internships", "internship experience"],
    "projects": ["projects", "personal projects", "academic projects", "technical projects", "key projects"],
    "skills": ["skills", "technical skills", "key skills", "core competencies", "skills and tools",
               "technical proficiency", "tools and technologies", "technologies"],
    "certifications": ["certifications", "certificates", "licenses", "certifications and licenses",
                       "licenses and certifications", "courses", "courses and certifications"],
}

# keyword -> section lookup built once
_HEADER_LOOKUP = {kw: section for section, kws in SECTION_HEADERS.items() for kw in kws}


def segment_sections(text: str) -> dict:
    """
    Split resume text into sections based on common header keywords.
    Returns a dict like {"education": "...", "experience": "...", ...}
    Anything before the first recognized header is treated as the 'contact' block.
    """
    lines = text.split("\n")
    sections = {"contact": []}
    current_section = "contact"

    for line in lines:
        matched_section = _match_section_header(line)
        if matched_section:
            current_section = matched_section
            sections.setdefault(current_section, [])
            continue
        sections.setdefault(current_section, []).append(line)

    return {key: "\n".join(value) for key, value in sections.items()}


def _match_section_header(line: str):
    """
    A header is a line that IS a section name (optionally with a trailing colon,
    decoration, or '&' for 'and') — not any short line that merely contains the word.
    Previously "Strong communication skills" or "Led 5 projects" were mistaken for headers.
    """
    normalized = line.strip().lower().replace("&", " and ")
    normalized = re.sub(r"[^a-z\s]", " ", normalized)       # drop colons, dashes, digits, bullets
    normalized = re.sub(r"\s+", " ", normalized).strip()
    return _HEADER_LOOKUP.get(normalized)
